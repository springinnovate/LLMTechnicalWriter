"""Pipeline to build a .json file to be later used the llm pipeline."""
import hashlib
import concurrent.futures
import pickle
import threading
from configparser import ConfigParser
import argparse
import datetime
import glob
import json
import logging
import os
import sys

from docx import Document
import llm_techincal_writer
import PyPDF2
import tiktoken
from openai import OpenAI

logging.basicConfig(
    level=logging.DEBUG,
    stream=sys.stdout,
    format=(
        '%(asctime)s (%(relativeCreated)d) %(levelname)s %(name)s'
        ' [%(funcName)s:%(lineno)d] %(message)s'))
logging.getLogger('openai').setLevel(logging.INFO)
logging.getLogger('httpcore').setLevel(logging.INFO)
logging.getLogger('httpx').setLevel(logging.WARNING)

LOGGER = logging.getLogger(__name__)

with open('pipeline_schema.json', 'r', encoding='utf-8') as f:
    PIPELINE_SCHEMA = json.loads(f.read())


def read_docx_to_text(docx_file_path):
    doc = Document(docx_file_path)
    text_parts = []
    for paragraph in doc.paragraphs:
        text_parts.append(paragraph.text)
    return "\n".join(text_parts)


def build_pipeline_with_placeholders(
        openai_context,
        form_description_text,
        openai_model):
    prompt_dict = {
        'developer': (
            'You are given an English description of a formal document that needs to be written. The description identifies its sections, lengths, and required content. Generate a structured pipeline based on these details. For each "file" in the preprocessing stage, the "prompt" must instruct how to parse or gather structured text needed in the analysis stage (e.g., which fields or data need to be extracted). Avoid prompts that are purely user instructions without producing text for analysis. The final output must map each relevant preprocessing section into the analysis and output steps.'),
        'user': (
            'Read the description along with the developer’s instructions. Identify '
            'the sections from the text and map them to any input files or data '
            'that the document implies, the analyses required, and the output '
            'format. If the text indicates additional data or context (e.g., '
            'personal details, references, or attachments), include them as '
            'required inputs, for a placeholder make any input files have the txt extension. Then produce a pipeline configuration that '
            'references these sections and satisfies the provided JSON schema. '
            'Wherever an analysis or output section must refer to content from a '
            'preprocessing section, include that section’s ID in curly braces '
            '(e.g., {appointments_positions}) in the assistant_template.'),
        'assistant': f'{form_description_text}\n\nJSONschema: {PIPELINE_SCHEMA}'
    }

    pipeline_text = llm_techincal_writer.generate_text(
        openai_context, prompt_dict, openai_model, force_regenerate=False)
    return pipeline_text


def create_template_from_schema(schema):
    schema_type = schema.get("type", None)

    if schema_type == "object":
        # Handle properties
        template = {}
        properties = schema.get("properties", {})
        required_props = schema.get("required", [])

        # Add explicit properties
        for prop_name, prop_schema in properties.items():
            template[prop_name] = create_template_from_schema(prop_schema) if prop_name in required_props else None

        # If there are patternProperties, add a single example pattern
        pattern_props = schema.get("patternProperties", {})
        for pattern, pattern_schema in pattern_props.items():
            # Use a generic key for demonstration (e.g. "example_key")
            template["example_key"] = create_template_from_schema(pattern_schema)

        return template

    elif schema_type == "array":
        # Handle array
        items_schema = schema.get("items", {})
        return [create_template_from_schema(items_schema)]

    elif schema_type == "string":
        return "string placeholder"

    # Fallback placeholder if no type
    return "placeholder"


def rewrite_pipeline_file_paths_and_create_files(
        pipeline_json, input_subdirectory="inputs"):
    """
    Reads a pipeline JSON (already loaded as a Python dict),
    then for every 'file_path' in the 'preprocessing' section:
      1) Updates 'file_path' to point to the new subdirectory (e.g. 'inputs/filename.txt').
      2) Ensures that subdirectory exists.
      3) Creates each file and writes the parent preprocessing 'description' into it.

    :param pipeline_json: The pipeline dict loaded from a JSON configuration.
    :param input_subdirectory: The subdirectory to place all input files.
    :return: The modified pipeline dict with updated file paths.
    """

    # Ensure the subdirectory exists
    os.makedirs(input_subdirectory, exist_ok=True)

    # Iterate through each preprocessing section
    preprocessing = pipeline_json.get("preprocessing", {})
    for section_id, section_data in preprocessing.items():
        # Grab the descriptive text we want to write into the file
        description_text = section_data.get("description", "")
        files_list = section_data.get("files", [])

        for file_obj in files_list:
            original_path = file_obj["file_path"]
            # Build new path inside the chosen subdirectory
            new_path = f"{input_subdirectory}/{os.path.basename(original_path)}"
            file_obj["file_path"] = new_path  # Update the pipeline to reference the new path

            # Create the file and write the description text
            with open(new_path, "w", encoding="utf-8") as f:
                f.write(description_text)

    return pipeline_json


def main():
    parser = argparse.ArgumentParser(
        description='Pipeline builder -- LLM Technical Writing Assistant')
    parser.add_argument(
        'form_description_path',
        help='Path to a docx describing the form to be generated.')

    parser.add_argument(
        '--openai_model',
        default='gpt-4o-mini',
        help='Select an openai model.')
    args = parser.parse_args()

    openai_context = llm_techincal_writer.create_openai_context()
    form_description_text = read_docx_to_text(args.form_description_path)
    form_basename = os.path.splitext(os.path.basename(args.form_description_path))[0]
    pipeline_text = build_pipeline_with_placeholders(
        openai_context,
        form_description_text,
        args.openai_model)
    intermediate_pipeline = json.loads(pipeline_text)
    final_pipeline_model = rewrite_pipeline_file_paths_and_create_files(
        intermediate_pipeline, input_subdirectory=form_basename)
    print(json.dumps(final_pipeline_model, indent=2))

if __name__ == '__main__':
    main()
